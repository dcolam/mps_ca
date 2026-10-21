"""Experiment metadata: which files belong to an experiment and what condition each covers.

The Excel workbook (one sheet per experiment) is only an import format. Columns are
located by header name, not position, and condition-level cells are forward-filled
the same way the original notebooks did.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

# canonical key -> accepted header spellings (lower-case, text before any "(")
_HEADER_ALIASES = {
    "main_folder": ("main folder", "main_folder", "experiment", "folder"),
    "subfolder": ("subfolder", "sub folder"),
    "filename": ("file name", "filename", "file"),
    "application": ("application", "condition"),
    "phase": ("phase",),
    "duration": ("time", "duration"),
    "frames": ("frames", "n frames"),
    "binning": ("binning", "camera binning"),
    "use_for_seg": ("segmentation", "seg", "use for segmentation"),
}


@dataclass
class FileEntry:
    filename: str
    application: str | None = None
    phase: str | None = None
    duration_s: int | None = None
    expected_frames: int | None = None  # source frames for the whole condition block
    block: int = 0                      # condition block index within the experiment
    use_for_seg: bool = False
    camera_binning: str | None = None   # as acquired, e.g. "2*2"
    subfolder: str | None = None
    path: Path | None = None


@dataclass
class ExperimentMeta:
    name: str                 # experiment folder name
    sheet: str | None = None
    genotype: str = "unknown"
    files: list[FileEntry] = field(default_factory=list)


def infer_genotype(*names: str | None) -> str:
    for n in names:
        if not n:
            continue
        if "KO" in n:
            return "KO"
        if "WT" in n:
            return "WT"
    return "unknown"


def _parse_duration(v) -> int | None:
    if v is None:
        return None
    try:
        return int(str(v).lower().replace("s", "").strip())
    except ValueError:
        return None


def _parse_int(v) -> int | None:
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


def _map_headers(header_row) -> dict[str, int]:
    cols = {}
    for i, h in enumerate(header_row):
        if h is None:
            continue
        key = str(h).split("(")[0].strip().lower()
        for canon, aliases in _HEADER_ALIASES.items():
            if key in aliases and canon not in cols:
                cols[canon] = i
    missing = {"main_folder", "filename"} - cols.keys()
    if missing:
        raise ValueError(f"missing required column(s): {sorted(missing)}")
    return cols


def parse_excel(xlsx_path: str | Path) -> list[ExperimentMeta]:
    """Parse one experiment per sheet from the metadata workbook."""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    experiments = []
    for sheet in wb.sheetnames:
        rows = list(wb[sheet].iter_rows(values_only=True))
        if not rows:
            continue
        try:
            cols = _map_headers(rows[0])
        except ValueError:
            continue

        def cell(row, key):
            i = cols.get(key)
            return row[i] if i is not None and i < len(row) else None

        exp_name = app = phase = dur = frames = sub = cam_bin = None
        block = -1
        files = []
        for row in rows[1:]:
            if cell(row, "main_folder"):
                exp_name = str(cell(row, "main_folder")).strip()
            # a new condition block starts whenever a condition-level cell is filled
            if any(cell(row, k) is not None for k in ("application", "phase", "frames")):
                block += 1
            if cell(row, "application"):
                app = str(cell(row, "application")).strip()
            if cell(row, "phase"):
                phase = str(cell(row, "phase")).strip()
            if cell(row, "duration"):
                dur = _parse_duration(cell(row, "duration"))
            if cell(row, "binning") is not None:
                cam_bin = str(cell(row, "binning")).strip()
            if cell(row, "frames") is not None:
                frames = _parse_int(cell(row, "frames"))
            s = cell(row, "subfolder")
            if s is not None:
                # some sheets hold frame counts in this column; only text names a folder
                sub = None if isinstance(s, (int, float)) else str(s).strip()
            fname = cell(row, "filename")
            if not fname:
                continue
            fname = str(fname).strip()
            if not fname.lower().endswith((".tif", ".tiff")):
                fname += ".tif"
            files.append(FileEntry(
                filename=fname, application=app, phase=phase, duration_s=dur,
                expected_frames=frames, block=max(block, 0),
                use_for_seg=bool(cell(row, "use_for_seg")), camera_binning=cam_bin,
                subfolder=sub,
            ))
        if exp_name:
            experiments.append(ExperimentMeta(
                name=exp_name, sheet=sheet,
                genotype=infer_genotype(sheet, exp_name), files=files,
            ))
    wb.close()
    return experiments


def _num_key(p: Path):
    return tuple(int(n) for n in re.findall(r"\d+", p.name))


def resolve_paths(exp_dir: str | Path, meta: ExperimentMeta) -> list[str]:
    """Fill in ``FileEntry.path`` by searching ``exp_dir``; returns the missing filenames.

    Files are matched by name anywhere under the folder; when the same name exists in
    several subfolders, the one in ``entry.subfolder`` wins.
    """
    exp_dir = Path(exp_dir)
    index: dict[str, list[Path]] = {}
    for p in sorted(exp_dir.rglob("*.tif*"), key=_num_key):
        index.setdefault(p.name, []).append(p)
    missing = []
    for fe in meta.files:
        cands = index.get(fe.filename, [])
        if fe.subfolder:
            in_sub = [p for p in cands if p.parent.name == fe.subfolder]
            cands = in_sub or cands
        if cands:
            fe.path = cands[0]
        else:
            fe.path = None
            missing.append(fe.filename)
    return missing


def entries_from_files(paths, application: str | None = None) -> list[FileEntry]:
    """Build entries for loose files without a workbook (one condition block per file)."""
    out = []
    for i, p in enumerate(paths):
        p = Path(p)
        out.append(FileEntry(filename=p.name, application=application or p.stem,
                             block=i, use_for_seg=True, path=p))
    return out


def _natural_key(p: Path):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", str(p))]


# folders like "40X" hold a single reference snapshot taken with another objective
_REFERENCE_DIR = re.compile(r"^\d+\s*[xX]$")


def entries_from_folder(folder: str | Path) -> ExperimentMeta:
    """Collect every TIFF under ``folder``; each subfolder becomes one condition block.

    Files are ordered naturally (``x_2`` before ``x_10``) so Micro-Manager split files
    stay in acquisition order. Objective reference folders (``40X``) are skipped.
    """
    folder = Path(folder)
    paths = sorted((p for p in folder.rglob("*.tif*") if p.is_file()
                    and not any(_REFERENCE_DIR.match(part) for part in p.relative_to(folder).parts[:-1])),
                   key=lambda p: _natural_key(p.relative_to(folder)))
    blocks: dict[Path, int] = {}
    files = []
    for p in paths:
        parent = p.parent.relative_to(folder)
        block = blocks.setdefault(parent, len(blocks))
        name = str(parent) if str(parent) != "." else p.stem
        files.append(FileEntry(filename=p.name, application=name, block=block,
                               use_for_seg=True, subfolder=p.parent.name, path=p))
    return ExperimentMeta(name=folder.name, genotype=infer_genotype(folder.name), files=files)
